from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Protocol

from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.app.api.errors import make_trace_id
from backend.app.core.config import Settings, get_settings
from backend.app.core.errors import NotFoundDomainError, ValidationDomainError
from backend.app.models import (
    AssessmentReport,
    Course,
    ExportJob,
    GeneratedResource,
    KnowledgePoint,
    LearningPath,
    LearningTask,
    PracticeAnswer,
    PracticeSession,
    User,
    WeaknessReviewItem,
)
from backend.app.schemas.exports import ExportJobResponse, LearningDossierExport, LearningDossierSourceSummary
from backend.app.schemas.reports import empty_report, iso_timestamp


class ExportNotFoundError(NotFoundDomainError):
    pass


class ExportValidationError(ValidationDomainError):
    pass


@dataclass(frozen=True)
class ExportUserRef:
    id: int


class ExportRepository(Protocol):
    def get_course_for_user(self, user_id: int, course_id: int) -> Course | None: ...

    def list_knowledge_points(self, course_id: int) -> list[KnowledgePoint]: ...

    def list_weakness_review_items(self, user_id: int, course_id: int) -> list[WeaknessReviewItem]: ...

    def get_active_path(self, user_id: int, course_id: int) -> LearningPath | None: ...

    def list_tasks_for_path(self, path_id: int) -> list[LearningTask]: ...

    def list_generated_resources(self, user_id: int, course_id: int) -> list[GeneratedResource]: ...

    def list_practice_answers(self, user_id: int, course_id: int) -> list[PracticeAnswer]: ...

    def get_latest_report(self, user_id: int, course_id: int) -> AssessmentReport | None: ...

    def get_generated_resource_for_user(self, user_id: int, resource_id: int) -> GeneratedResource | None: ...

    def list_export_jobs_for_resource(self, user_id: int, resource_id: int) -> list[ExportJob]: ...

    def add_export_job(self, job: ExportJob) -> ExportJob: ...

    def get_export_job_for_user(self, user_id: int, job_id: int) -> ExportJob | None: ...

    def get_export_job(self, job_id: int) -> ExportJob | None: ...

    def commit(self) -> None: ...

    def rollback(self) -> None: ...

    def refresh(self, instance: object) -> None: ...


class ExportJobQueue(Protocol):
    def enqueue(self, job_id: int) -> None: ...


@dataclass(frozen=True)
class RenderedExport:
    filename: str
    content_type: str
    content: bytes


class RqExportJobQueue:
    def __init__(self, redis_url: str, queue_name: str) -> None:
        self.redis_url = redis_url
        self.queue_name = queue_name

    def enqueue(self, job_id: int) -> None:
        from redis import Redis
        from rq import Queue

        from backend.app.workers.export_jobs import run_export_job

        connection = Redis.from_url(self.redis_url)
        queue = Queue(self.queue_name, connection=connection)
        queue.enqueue(run_export_job, job_id)


class SqlAlchemyExportRepository:
    def __init__(self, db: Session) -> None:
        self.db = db

    def get_course_for_user(self, user_id: int, course_id: int) -> Course | None:
        return self.db.scalar(select(Course).where(Course.id == course_id, Course.owner_id == user_id))

    def list_knowledge_points(self, course_id: int) -> list[KnowledgePoint]:
        return list(
            self.db.scalars(
                select(KnowledgePoint).where(KnowledgePoint.course_id == course_id).order_by(KnowledgePoint.order_index, KnowledgePoint.id)
            )
        )

    def list_weakness_review_items(self, user_id: int, course_id: int) -> list[WeaknessReviewItem]:
        return list(
            self.db.scalars(
                select(WeaknessReviewItem)
                .where(WeaknessReviewItem.user_id == user_id, WeaknessReviewItem.course_id == course_id)
                .order_by(WeaknessReviewItem.updated_at.desc(), WeaknessReviewItem.id.desc())
            )
        )

    def get_active_path(self, user_id: int, course_id: int) -> LearningPath | None:
        return self.db.scalar(
            select(LearningPath)
            .where(LearningPath.user_id == user_id, LearningPath.course_id == course_id, LearningPath.status == "active")
            .order_by(LearningPath.updated_at.desc(), LearningPath.id.desc())
        )

    def list_tasks_for_path(self, path_id: int) -> list[LearningTask]:
        return list(
            self.db.scalars(select(LearningTask).where(LearningTask.path_id == path_id).order_by(LearningTask.id.asc()))
        )

    def list_generated_resources(self, user_id: int, course_id: int) -> list[GeneratedResource]:
        return list(
            self.db.scalars(
                select(GeneratedResource)
                .where(GeneratedResource.user_id == user_id, GeneratedResource.course_id == course_id)
                .order_by(GeneratedResource.updated_at.desc(), GeneratedResource.id.desc())
            )
        )

    def list_practice_answers(self, user_id: int, course_id: int) -> list[PracticeAnswer]:
        session_ids = select(PracticeSession.id).where(
            PracticeSession.user_id == user_id,
            PracticeSession.course_id == course_id,
            PracticeSession.status == "completed",
        )
        return list(
            self.db.scalars(
                select(PracticeAnswer)
                .where(PracticeAnswer.user_id == user_id, PracticeAnswer.session_id.in_(session_ids))
                .order_by(PracticeAnswer.created_at.desc(), PracticeAnswer.id.desc())
            )
        )

    def get_latest_report(self, user_id: int, course_id: int) -> AssessmentReport | None:
        return self.db.scalar(
            select(AssessmentReport)
            .where(AssessmentReport.user_id == user_id, AssessmentReport.course_id == course_id)
            .order_by(AssessmentReport.created_at.desc(), AssessmentReport.id.desc())
        )

    def get_generated_resource_for_user(self, user_id: int, resource_id: int) -> GeneratedResource | None:
        return self.db.scalar(
            select(GeneratedResource).where(
                GeneratedResource.id == resource_id,
                GeneratedResource.user_id == user_id,
            )
        )

    def list_export_jobs_for_resource(self, user_id: int, resource_id: int) -> list[ExportJob]:
        return list(
            self.db.scalars(
                select(ExportJob)
                .where(ExportJob.user_id == user_id, ExportJob.resource_id == resource_id)
                .order_by(ExportJob.created_at.desc(), ExportJob.id.desc())
            )
        )

    def add_export_job(self, job: ExportJob) -> ExportJob:
        self.db.add(job)
        self.db.flush()
        return job

    def get_export_job_for_user(self, user_id: int, job_id: int) -> ExportJob | None:
        return self.db.scalar(select(ExportJob).where(ExportJob.id == job_id, ExportJob.user_id == user_id))

    def get_export_job(self, job_id: int) -> ExportJob | None:
        return self.db.scalar(select(ExportJob).where(ExportJob.id == job_id))

    def commit(self) -> None:
        self.db.commit()

    def rollback(self) -> None:
        self.db.rollback()

    def refresh(self, instance: object) -> None:
        self.db.refresh(instance)


class ExportService:
    allowed_formats = {"markdown", "pdf", "docx"}
    resource_formats = {"pptx"}

    def __init__(
        self,
        repository: ExportRepository,
        settings: Settings | None = None,
        job_queue: ExportJobQueue | None = None,
        run_jobs_inline: bool = False,
    ) -> None:
        self.repository = repository
        self.settings = settings or get_settings()
        self.job_queue = job_queue
        self.run_jobs_inline = run_jobs_inline

    def export_learning_dossier(self, user: User, course_id: int) -> LearningDossierExport:
        course = self.repository.get_course_for_user(user.id, course_id)
        if course is None:
            raise ExportNotFoundError("课程不存在或无权访问。")

        generated_at = datetime.now(UTC)
        points = self.repository.list_knowledge_points(course.id)
        weaknesses = self.repository.list_weakness_review_items(user.id, course.id)
        active_path = self.repository.get_active_path(user.id, course.id)
        tasks = self.repository.list_tasks_for_path(active_path.id) if active_path is not None else []
        resources = self.repository.list_generated_resources(user.id, course.id)
        answers = self.repository.list_practice_answers(user.id, course.id)
        report = self.repository.get_latest_report(user.id, course.id)
        report_body = report.report_json if report is not None else empty_report(course.id).report
        agent_trace_id = make_trace_id()

        source_summary = LearningDossierSourceSummary(
            has_report=report is not None,
            report_id=str(report.id) if report is not None else None,
            knowledge_point_count=len(points),
            weakness_count=len([item for item in weaknesses if item.status != "dismissed"]),
            path_task_count=len(tasks),
            resource_count=len(resources),
            practice_answer_count=len(answers),
        )
        markdown = self._build_markdown(
            course=course,
            report=report,
            report_body=report_body,
            points=points,
            weaknesses=weaknesses,
            active_path=active_path,
            tasks=tasks,
            resources=resources,
            answers=answers,
            generated_at=generated_at,
            source_summary=source_summary,
        )
        return LearningDossierExport(
            course_id=str(course.id),
            agent_trace_id=agent_trace_id,
            filename=self._safe_filename(course.title),
            content_type="text/markdown; charset=utf-8",
            markdown=markdown,
            generated_at=iso_timestamp(generated_at) or "",
            source_summary=source_summary,
        )

    def create_learning_dossier_job(self, user: User, course_id: int, export_format: str = "markdown") -> ExportJobResponse:
        normalized_format = self._normalize_format(export_format)
        course = self.repository.get_course_for_user(user.id, course_id)
        if course is None:
            raise ExportNotFoundError("课程不存在或无权访问。")

        job = ExportJob(
            user_id=user.id,
            course_id=course.id,
            export_type="learning_dossier",
            export_format=normalized_format,
            status="queued",
            filename=None,
            content_type=None,
            file_path=None,
            error_message=None,
            agent_trace_id=make_trace_id(),
            metadata_json={"course_title": self._text(course.title), "source": "learning_dossier"},
        )

        try:
            self.repository.add_export_job(job)
            self.repository.commit()
            self.repository.refresh(job)
        except Exception:
            self.repository.rollback()
            raise

        if self.run_jobs_inline:
            self.run_export_job(int(job.id))
            refreshed = self.repository.get_export_job_for_user(user.id, int(job.id)) or job
            return self._job_response(refreshed)

        if self.job_queue is not None:
            try:
                self.job_queue.enqueue(int(job.id))
            except Exception:
                self._mark_job_failed(job, "导出任务排队失败，请稍后重试。")
                return self._job_response(job)
        return self._job_response(job)

    def create_resource_export_job(self, user: User, resource_id: int, export_format: str = "pptx") -> ExportJobResponse:
        normalized_format = export_format.strip().lower()
        if normalized_format not in self.resource_formats:
            raise ExportValidationError("资源导出格式只能是 pptx。")
        resource = self.repository.get_generated_resource_for_user(user.id, resource_id)
        if resource is None:
            raise ExportNotFoundError("资源不存在或无权访问。")
        if resource.resource_type != "slide":
            raise ExportValidationError("只有 PPT 资源可以导出为 pptx。")

        existing_jobs = self.repository.list_export_jobs_for_resource(user.id, resource.id)
        reusable = next(
            (
                job
                for job in existing_jobs
                if job.export_format == normalized_format and job.status in {"queued", "running", "completed"}
            ),
            None,
        )
        if reusable is not None:
            return self._job_response(reusable)

        job = ExportJob(
            user_id=user.id,
            course_id=resource.course_id,
            resource_id=resource.id,
            export_type="resource_artifact",
            export_format=normalized_format,
            status="queued",
            filename=None,
            content_type=None,
            file_path=None,
            error_message=None,
            agent_trace_id=resource.agent_trace_id or make_trace_id(),
            metadata_json={"resource_type": "slide", "resource_title": self._text(resource.title)},
        )
        try:
            self.repository.add_export_job(job)
            self.repository.commit()
            self.repository.refresh(job)
        except Exception:
            self.repository.rollback()
            raise

        if self.run_jobs_inline:
            self.run_export_job(int(job.id))
            refreshed = self.repository.get_export_job_for_user(user.id, int(job.id)) or job
            return self._job_response(refreshed)
        if self.job_queue is not None:
            try:
                self.job_queue.enqueue(int(job.id))
            except Exception:
                self._mark_job_failed(job, "PPTX 导出任务排队失败，请稍后重试。")
        return self._job_response(job)

    def list_resource_export_jobs(self, user: User, resource_id: int) -> list[ExportJobResponse]:
        resource = self.repository.get_generated_resource_for_user(user.id, resource_id)
        if resource is None:
            raise ExportNotFoundError("资源不存在或无权访问。")
        return [self._job_response(job) for job in self.repository.list_export_jobs_for_resource(user.id, resource.id)]

    def run_export_job(self, job_id: int) -> ExportJobResponse:
        job = self.repository.get_export_job(job_id)
        if job is None:
            raise ExportNotFoundError("导出任务不存在。")
        if job.status == "completed":
            return self._job_response(job)

        try:
            job.status = "running"
            job.updated_at = datetime.now(UTC)
            self.repository.commit()
            rendered = self._render_export_job(job)
            export_dir = Path(self.settings.export_dir)
            export_dir.mkdir(parents=True, exist_ok=True)
            file_path = export_dir / f"job-{job.id}-{rendered.filename}"
            file_path.write_bytes(rendered.content)
            job.status = "completed"
            job.filename = rendered.filename
            job.content_type = rendered.content_type
            job.file_path = str(file_path)
            job.error_message = None
            job.completed_at = datetime.now(UTC)
            job.updated_at = job.completed_at
            self.repository.commit()
            self.repository.refresh(job)
        except Exception:
            self.repository.rollback()
            failed_job = self.repository.get_export_job(job_id)
            if failed_job is not None:
                self._mark_job_failed(failed_job, "学习档案导出失败，请稍后重试。")
                job = failed_job
        return self._job_response(job)

    def _render_export_job(self, job: ExportJob) -> RenderedExport:
        if job.export_type == "resource_artifact":
            return self._render_resource_artifact_job(job)
        return self._render_learning_dossier_job(job)

    def get_export_job(self, user: User, job_id: int) -> ExportJobResponse:
        job = self.repository.get_export_job_for_user(user.id, job_id)
        if job is None:
            raise ExportNotFoundError("导出任务不存在或无权访问。")
        return self._job_response(job)

    def get_export_job_file(self, user: User, job_id: int) -> ExportJob:
        job = self.repository.get_export_job_for_user(user.id, job_id)
        if job is None:
            raise ExportNotFoundError("导出任务不存在或无权访问。")
        if job.status != "completed" or not job.file_path:
            raise ExportNotFoundError("导出任务尚未完成。")
        return job

    def _render_learning_dossier_job(self, job: ExportJob) -> RenderedExport:
        if job.course_id is None:
            raise ExportNotFoundError("导出任务缺少课程。")
        dossier = self.export_learning_dossier(ExportUserRef(int(job.user_id)), int(job.course_id))  # type: ignore[arg-type]
        export_format = self._normalize_format(job.export_format)
        base_filename = self._replace_suffix(dossier.filename, ".md")
        if export_format == "markdown":
            return RenderedExport(
                filename=f"{base_filename}.md",
                content_type="text/markdown; charset=utf-8",
                content=dossier.markdown.encode("utf-8"),
            )
        if export_format == "pdf":
            return RenderedExport(
                filename=f"{base_filename}.pdf",
                content_type="application/pdf",
                content=self._render_pdf(dossier.markdown),
            )
        return RenderedExport(
            filename=f"{base_filename}.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            content=self._render_docx(dossier.markdown),
        )

    def _render_resource_artifact_job(self, job: ExportJob) -> RenderedExport:
        if job.resource_id is None:
            raise ExportNotFoundError("资源导出任务缺少资源。")
        resource = self.repository.get_generated_resource_for_user(int(job.user_id), int(job.resource_id))
        if resource is None:
            raise ExportNotFoundError("资源不存在或无权访问。")
        if resource.resource_type != "slide" or job.export_format != "pptx":
            raise ExportValidationError("当前资源不支持该导出格式。")
        slides = self._resource_slides(resource)
        filename = self._safe_resource_filename(resource.title, ".pptx")
        return RenderedExport(
            filename=filename,
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            content=self._render_pptx(resource.title, slides, list(resource.citation_json or [])),
        )

    def _render_pptx(self, title: str, slides: list[dict[str, Any]], citations: list[dict[str, Any]]) -> bytes:
        from io import BytesIO

        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt

        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        blank_layout = presentation.slide_layouts[6]
        accent = RGBColor(15, 143, 131)
        ink = RGBColor(16, 27, 35)
        muted = RGBColor(93, 112, 111)
        surface = RGBColor(246, 250, 249)

        for index, slide_data in enumerate(slides, start=1):
            slide = presentation.slides.add_slide(blank_layout)
            background = slide.background.fill
            background.solid()
            background.fore_color.rgb = surface

            marker = slide.shapes.add_shape(1, Inches(0.62), Inches(0.55), Inches(0.12), Inches(0.72))
            marker.fill.solid()
            marker.fill.fore_color.rgb = accent
            marker.line.fill.background()

            title_box = slide.shapes.add_textbox(Inches(0.95), Inches(0.48), Inches(11.5), Inches(0.9))
            title_frame = title_box.text_frame
            title_frame.clear()
            title_paragraph = title_frame.paragraphs[0]
            title_paragraph.text = self._text(slide_data.get("title") or title)
            title_paragraph.font.name = "Microsoft YaHei"
            title_paragraph.font.size = Pt(28 if index > 1 else 32)
            title_paragraph.font.bold = True
            title_paragraph.font.color.rgb = ink

            body_box = slide.shapes.add_textbox(Inches(0.98), Inches(1.65), Inches(11.2), Inches(3.75))
            body_frame = body_box.text_frame
            body_frame.clear()
            body_frame.word_wrap = True
            bullets = slide_data.get("bullets") if isinstance(slide_data.get("bullets"), list) else []
            for bullet_index, bullet in enumerate(bullets[:6]):
                paragraph = body_frame.paragraphs[0] if bullet_index == 0 else body_frame.add_paragraph()
                paragraph.text = self._text(bullet)[:240]
                paragraph.level = 0
                paragraph.space_after = Pt(13)
                paragraph.font.name = "Microsoft YaHei"
                paragraph.font.size = Pt(20)
                paragraph.font.color.rgb = ink

            notes_box = slide.shapes.add_textbox(Inches(0.98), Inches(5.65), Inches(11.2), Inches(0.9))
            notes_frame = notes_box.text_frame
            notes_frame.clear()
            notes_paragraph = notes_frame.paragraphs[0]
            notes_paragraph.text = f"讲稿提示：{self._text(slide_data.get('speaker_notes'))[:360]}"
            notes_paragraph.font.name = "Microsoft YaHei"
            notes_paragraph.font.size = Pt(11)
            notes_paragraph.font.color.rgb = muted

            footer = slide.shapes.add_textbox(Inches(0.98), Inches(6.85), Inches(11.2), Inches(0.28))
            footer_paragraph = footer.text_frame.paragraphs[0]
            footer_paragraph.text = f"EduNova · {index}/{len(slides) + (1 if citations else 0)}"
            footer_paragraph.alignment = PP_ALIGN.RIGHT
            footer_paragraph.font.name = "Microsoft YaHei"
            footer_paragraph.font.size = Pt(9)
            footer_paragraph.font.color.rgb = muted

        if citations:
            citation_slide = presentation.slides.add_slide(blank_layout)
            citation_slide.background.fill.solid()
            citation_slide.background.fill.fore_color.rgb = surface
            title_box = citation_slide.shapes.add_textbox(Inches(0.95), Inches(0.65), Inches(11.2), Inches(0.7))
            paragraph = title_box.text_frame.paragraphs[0]
            paragraph.text = "课程引用"
            paragraph.font.name = "Microsoft YaHei"
            paragraph.font.size = Pt(28)
            paragraph.font.bold = True
            paragraph.font.color.rgb = ink
            body = citation_slide.shapes.add_textbox(Inches(0.98), Inches(1.65), Inches(11.1), Inches(4.8))
            frame = body.text_frame
            frame.clear()
            for citation_index, citation in enumerate(citations[:8], start=1):
                item = frame.paragraphs[0] if citation_index == 1 else frame.add_paragraph()
                section = self._text(citation.get("section_title") or "课程章节")
                source = self._text(citation.get("source_title") or "课程资料")
                item.text = f"{citation_index}. {section}（{source}）"
                item.font.name = "Microsoft YaHei"
                item.font.size = Pt(17)
                item.font.color.rgb = ink
                item.space_after = Pt(10)

        buffer = BytesIO()
        presentation.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _resource_slides(resource: GeneratedResource) -> list[dict[str, Any]]:
        content = resource.content_json or {}
        artifact = content.get("artifact")
        if isinstance(artifact, dict) and artifact.get("kind") == "slide_deck":
            slides = artifact.get("slides")
            if isinstance(slides, list) and slides:
                return [slide for slide in slides if isinstance(slide, dict)]

        markdown = str(content.get("markdown") or "")
        matches = list(re.finditer(r"^## 第 \d+ 页：(.+)$", markdown, flags=re.MULTILINE))
        slides: list[dict[str, Any]] = []
        for index, match in enumerate(matches):
            end = matches[index + 1].start() if index + 1 < len(matches) else len(markdown)
            block = markdown[match.end() : end]
            bullets = [line.removeprefix("- 要点：").strip() for line in block.splitlines() if line.startswith("- 要点：")]
            notes = next((line.removeprefix("讲稿：").strip() for line in block.splitlines() if line.startswith("讲稿：")), "")
            slides.append({"title": match.group(1).strip(), "bullets": bullets or ["课程内容"], "speaker_notes": notes})
        if not slides:
            raise ExportValidationError("PPT 资源缺少可导出的页面结构。")
        return slides

    def _render_pdf(self, markdown: str) -> bytes:
        from io import BytesIO

        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        font_name = self._resolve_pdf_font()
        buffer = BytesIO()
        pdf = canvas.Canvas(buffer, pagesize=A4)
        width, height = A4
        left = 48
        y = height - 52
        pdf.setFont(font_name, 11)
        for raw_line in markdown.splitlines():
            line = raw_line.strip() or " "
            for segment in self._wrap_text(line, 48):
                if y < 48:
                    pdf.showPage()
                    pdf.setFont(font_name, 11)
                    y = height - 52
                pdf.drawString(left, y, segment)
                y -= 17
        pdf.save()
        return buffer.getvalue()

    def _resolve_pdf_font(self) -> str:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfbase.ttfonts import TTFont

        font_name = "EduNovaCJK"
        font_path = self._find_cjk_font()
        if font_path is not None:
            try:
                if font_name not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
                return font_name
            except Exception:
                pass

        cid_font_name = "STSong-Light"
        if cid_font_name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(UnicodeCIDFont(cid_font_name))
        return cid_font_name

    def _render_docx(self, markdown: str) -> bytes:
        from io import BytesIO

        from docx import Document

        document = Document()
        for line in markdown.splitlines():
            stripped = line.strip()
            if not stripped:
                document.add_paragraph("")
            elif stripped.startswith("# "):
                document.add_heading(stripped[2:].strip(), level=1)
            elif stripped.startswith("## "):
                document.add_heading(stripped[3:].strip(), level=2)
            else:
                document.add_paragraph(stripped)
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def _mark_job_failed(self, job: ExportJob, message: str) -> None:
        try:
            job.status = "failed"
            job.error_message = message
            job.updated_at = datetime.now(UTC)
            self.repository.commit()
            self.repository.refresh(job)
        except Exception:
            self.repository.rollback()
            raise

    def _job_response(self, job: ExportJob) -> ExportJobResponse:
        return ExportJobResponse(
            job_id=str(job.id),
            status=job.status,
            format=job.export_format,
            export_type=job.export_type,
            resource_id=str(job.resource_id) if job.resource_id is not None else None,
            filename=job.filename,
            content_type=job.content_type,
            agent_trace_id=job.agent_trace_id,
            error_message=job.error_message,
            created_at=iso_timestamp(job.created_at) or "",
            updated_at=iso_timestamp(job.updated_at) or "",
            completed_at=iso_timestamp(job.completed_at),
        )

    def _normalize_format(self, export_format: str) -> str:
        normalized = export_format.strip().lower()
        if normalized == "md":
            normalized = "markdown"
        if normalized not in self.allowed_formats:
            raise ExportValidationError("导出格式只能是 markdown、pdf 或 docx。")
        return normalized

    @staticmethod
    def _replace_suffix(filename: str, suffix: str) -> str:
        return filename[: -len(suffix)] if filename.endswith(suffix) else filename

    @staticmethod
    def _wrap_text(text: str, width: int) -> list[str]:
        if len(text) <= width:
            return [text]
        return [text[index : index + width] for index in range(0, len(text), width)]

    @staticmethod
    def _find_cjk_font() -> Path | None:
        candidates = [
            Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
            Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
            Path("C:/Windows/Fonts/msyh.ttc"),
            Path("C:/Windows/Fonts/simhei.ttf"),
        ]
        return next((path for path in candidates if path.exists()), None)

    def _build_markdown(
        self,
        course: Course,
        report: AssessmentReport | None,
        report_body: dict,
        points: list[KnowledgePoint],
        weaknesses: list[WeaknessReviewItem],
        active_path: LearningPath | None,
        tasks: list[LearningTask],
        resources: list[GeneratedResource],
        answers: list[PracticeAnswer],
        generated_at: datetime,
        source_summary: LearningDossierSourceSummary,
    ) -> str:
        lines: list[str] = [
            f"# {self._text(course.title)} 学习档案",
            "",
            f"- 导出时间：{iso_timestamp(generated_at) or ''}",
            f"- 课程：{self._text(course.title)}",
            f"- 学科：{self._text(course.subject or '未设置')}",
            f"- 资料来源：{self._text(course.source_type)}",
            "",
            "## 数据概览",
            "",
            f"- 知识点：{source_summary.knowledge_point_count} 个",
            f"- 弱点队列：{source_summary.weakness_count} 项",
            f"- 学习路径任务：{source_summary.path_task_count} 项",
            f"- 已生成资源：{source_summary.resource_count} 个",
            f"- 练习证据：{source_summary.practice_answer_count} 条",
            "",
            "## 学习报告",
            "",
        ]

        if report is None:
            lines.extend(["还没有真实学习报告。", "", "- 完成一次课程练习后生成报告。"])
        else:
            score = int(report.score) if report.score is not None else "暂无"
            lines.extend([f"- 报告 ID：{report.id}", f"- 报告时间：{iso_timestamp(report.created_at) or ''}", f"- 报告分数：{score}"])
            summary = self._text(str(report_body.get("summary") or "暂无报告摘要。"), limit=400)
            lines.extend(["", summary])

        mastery = report_body.get("mastery_update") if isinstance(report_body.get("mastery_update"), dict) else {}
        lines.extend(
            [
                "",
                "## 掌握度摘要",
                "",
                f"- 薄弱点：{self._int_value(mastery.get('weak_count'))}",
                f"- 已掌握：{self._int_value(mastery.get('mastered_count'))}",
                f"- 学习中：{self._int_value(mastery.get('learning_count'))}",
                "",
                "## 薄弱点与复习队列",
                "",
            ]
        )
        active_weaknesses = [item for item in weaknesses if item.status != "dismissed"]
        lines.extend(self._weakness_lines(active_weaknesses))
        lines.extend(["", "## 当前学习路径", ""])
        lines.extend(self._path_lines(active_path, tasks))
        lines.extend(["", "## 推荐资源", ""])
        lines.extend(self._resource_lines(resources))
        lines.extend(["", "## 练习证据摘要", ""])
        lines.extend(self._practice_lines(answers))
        lines.extend(["", "## 下一步建议", ""])
        lines.extend(self._suggestion_lines(report_body))
        lines.extend(["", "## 安全说明", "", "- 本档案只导出课程级学习摘要，不包含原始资料全文、内部指令、模型请求内容、密钥、登录令牌或完整用户画像。"])
        return "\n".join(lines).strip() + "\n"

    def _weakness_lines(self, weaknesses: list[WeaknessReviewItem]) -> list[str]:
        if not weaknesses:
            return ["还没有确认或待复习的弱点。"]
        return [
            (
                f"- {self._text(item.title)}"
                f"（状态：{self._text(item.status)}；来源：{self._text(item.source_type)}；"
                f"下次复习：{iso_timestamp(item.next_review_at) or '未设置'}）"
            )
            for item in weaknesses[:10]
        ]

    def _path_lines(self, active_path: LearningPath | None, tasks: list[LearningTask]) -> list[str]:
        if active_path is None:
            return ["还没有生成课程学习路径。"]
        lines = [f"- 路径：{self._text(active_path.title)}", f"- 目标：{self._text(active_path.goal or '未设置')}"]
        if not tasks:
            lines.append("- 暂无路径任务。")
            return lines
        lines.append("")
        lines.extend(
            [
                f"- [{self._text(task.status)}] {self._text(task.title)}"
                f"（类型：{self._text(task.task_type)}）"
                for task in tasks[:10]
            ]
        )
        return lines

    def _resource_lines(self, resources: list[GeneratedResource]) -> list[str]:
        if not resources:
            return ["还没有生成可推荐的课程资源。"]
        return [
            (
                f"- {self._text(resource.title)}"
                f"（类型：{self._text(resource.resource_type)}；审核：{self._text(resource.review_status)}；"
                f"置信度：{float(resource.confidence_score) if resource.confidence_score is not None else '暂无'}）"
            )
            for resource in resources[:10]
        ]

    def _practice_lines(self, answers: list[PracticeAnswer]) -> list[str]:
        if not answers:
            return ["还没有可用于导出的练习作答证据。"]
        scored = [int((answer.feedback_json or {}).get("score") or 0) for answer in answers]
        low_score_count = sum(1 for score in scored if score < 60)
        correct_count = sum(1 for answer in answers if answer.is_correct is True)
        return [
            f"- 已记录作答：{len(answers)} 条",
            f"- 正确作答：{correct_count} 条",
            f"- 低分或错误：{low_score_count} 条",
        ]

    def _suggestion_lines(self, report_body: dict) -> list[str]:
        suggestions = report_body.get("next_step_suggestions")
        if not isinstance(suggestions, list) or not suggestions:
            return ["完成一次课程练习后生成报告。"]
        return [f"- {self._text(str(item), limit=240)}" for item in suggestions[:8]]

    @staticmethod
    def _safe_filename(course_title: str) -> str:
        title = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", course_title).strip() or "course"
        title = re.sub(r"\s+", "-", title)[:80]
        return f"edunova-{title}-learning-dossier.md"

    @staticmethod
    def _safe_resource_filename(resource_title: str, suffix: str) -> str:
        title = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", resource_title).strip() or "resource"
        title = re.sub(r"\s+", "-", title)[:80]
        return f"edunova-{title}{suffix}"

    @staticmethod
    def _text(value: str, limit: int = 160) -> str:
        cleaned = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
        cleaned = " ".join(cleaned.split())
        if len(cleaned) > limit:
            return f"{cleaned[:limit].rstrip()}..."
        return cleaned

    @staticmethod
    def _int_value(value: object) -> int:
        try:
            return int(value)  # type: ignore[arg-type]
        except (TypeError, ValueError):
            return 0
