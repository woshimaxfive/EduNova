from __future__ import annotations

import argparse
from io import BytesIO
import json
from pathlib import Path
from time import perf_counter

from docx import Document
from pptx import Presentation
from reportlab.pdfgen import canvas

from backend.app.services.material_parsers import DoclingDocumentExtractor


def make_pdf() -> bytes:
    output = BytesIO()
    pdf = canvas.Canvas(output)
    pdf.drawString(72, 760, "Chapter 1 Linear Structures")
    pdf.drawString(72, 730, "A stack follows last in, first out semantics.")
    pdf.showPage()
    pdf.drawString(72, 760, "Chapter 2 Trees")
    pdf.drawString(72, 730, "A binary tree has at most two children per node.")
    pdf.save()
    return output.getvalue()


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Chapter 1 Linear Structures", level=1)
    document.add_paragraph("A queue follows first in, first out semantics.")
    document.add_heading("1.1 Stack", level=2)
    document.add_paragraph("A stack supports push and pop operations.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_pptx() -> bytes:
    presentation = Presentation()
    for title, body in (
        ("Chapter 1 Graphs", "Breadth-first search explores vertices by level."),
        ("Chapter 2 Sorting", "Merge sort uses divide and conquer."),
    ):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def benchmark(
    extractor: DoclingDocumentExtractor,
    *,
    name: str,
    extension: str,
    content: bytes,
) -> dict[str, int | str]:
    started = perf_counter()
    document = extractor.parse_document(extension, content)
    duration_ms = round((perf_counter() - started) * 1000)
    return {
        "name": name,
        "extension": extension,
        "duration_ms": duration_ms,
        "page_count": len(document.pages),
        "block_count": len(document.blocks),
        "heading_count": sum(block.kind == "heading" for block in document.blocks),
        "character_count": len(document.text),
        "parser": document.parser,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="运行 Docling 脱敏迁移基准，只输出聚合指标。")
    parser.add_argument("--artifacts-path", default="/opt/docling/models")
    parser.add_argument("--local-document", type=Path)
    args = parser.parse_args()

    extractor = DoclingDocumentExtractor(artifacts_path=args.artifacts_path, timeout_seconds=180)
    fixtures = {
        ".pdf": make_pdf(),
        ".docx": make_docx(),
        ".pptx": make_pptx(),
    }
    results = [
        benchmark(extractor, name=f"synthetic{extension}", extension=extension, content=content)
        for extension, content in fixtures.items()
    ]
    if args.local_document:
        local_path = args.local_document.resolve(strict=True)
        extension = local_path.suffix.casefold()
        if extension not in extractor.supported_extensions:
            raise SystemExit("本地基准只支持 PDF、DOCX 或 PPTX。")
        results.append(
            benchmark(
                extractor,
                name="local-redacted",
                extension=extension,
                content=local_path.read_bytes(),
            )
        )
    print(json.dumps({"fixtures": results}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
