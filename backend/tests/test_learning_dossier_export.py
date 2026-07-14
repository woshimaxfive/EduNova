from __future__ import annotations

from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient

from backend.app.api.v1.deps import get_auth_service
from backend.app.core.config import Settings
from backend.app.core.security import create_access_token
from backend.app.main import create_app
from backend.app.models import (
    AssessmentReport,
    Course,
    GeneratedResource,
    KnowledgePoint,
    LearningPath,
    LearningTask,
    PracticeAnswer,
    PracticeSession,
    User,
    WeaknessReviewItem,
)
from backend.app.services.auth import AuthService


NOW = datetime(2026, 7, 5, 15, 0, tzinfo=UTC)


@dataclass
class TokenAuthRepository:
    user: User

    def get_user_by_id(self, user_id: int) -> User | None:
        return self.user if self.user.id == user_id else None


@dataclass
class FakeLearningDossierRepository:
    courses: list[Course] = field(default_factory=list)
    knowledge_points: list[KnowledgePoint] = field(default_factory=list)
    weakness_items: list[WeaknessReviewItem] = field(default_factory=list)
    paths: list[LearningPath] = field(default_factory=list)
    tasks: list[LearningTask] = field(default_factory=list)
    resources: list[GeneratedResource] = field(default_factory=list)
    practice_sessions: list[PracticeSession] = field(default_factory=list)
    practice_answers: list[PracticeAnswer] = field(default_factory=list)
    reports: list[AssessmentReport] = field(default_factory=list)
    export_jobs: list[Any] = field(default_factory=list)
    next_export_job_id: int = 1

    def get_course_for_user(self, user_id: int, course_id: int) -> Course | None:
        return next((course for course in self.courses if course.id == course_id and course.owner_id == user_id), None)

    def list_knowledge_points(self, course_id: int) -> list[KnowledgePoint]:
        return sorted(
            [point for point in self.knowledge_points if point.course_id == course_id],
            key=lambda point: (point.order_index, point.id),
        )

    def list_weakness_review_items(self, user_id: int, course_id: int) -> list[WeaknessReviewItem]:
        return sorted(
            [item for item in self.weakness_items if item.user_id == user_id and item.course_id == course_id],
            key=lambda item: (item.updated_at, item.id),
            reverse=True,
        )

    def get_active_path(self, user_id: int, course_id: int) -> LearningPath | None:
        paths = [
            path
            for path in self.paths
            if path.user_id == user_id and path.course_id == course_id and path.status == "active"
        ]
        return sorted(paths, key=lambda path: (path.updated_at, path.id), reverse=True)[0] if paths else None

    def list_tasks_for_path(self, path_id: int) -> list[LearningTask]:
        return sorted([task for task in self.tasks if task.path_id == path_id], key=lambda task: task.id)

    def list_generated_resources(self, user_id: int, course_id: int) -> list[GeneratedResource]:
        return sorted(
            [resource for resource in self.resources if resource.user_id == user_id and resource.course_id == course_id],
            key=lambda resource: (resource.updated_at, resource.id),
            reverse=True,
        )

    def list_practice_answers(self, user_id: int, course_id: int) -> list[PracticeAnswer]:
        session_ids = {
            session.id
            for session in self.practice_sessions
            if session.user_id == user_id and session.course_id == course_id and session.status == "completed"
        }
        return [answer for answer in self.practice_answers if answer.user_id == user_id and answer.session_id in session_ids]

    def get_latest_report(self, user_id: int, course_id: int) -> AssessmentReport | None:
        reports = [report for report in self.reports if report.user_id == user_id and report.course_id == course_id]
        return sorted(reports, key=lambda report: (report.created_at, report.id), reverse=True)[0] if reports else None

    def get_generated_resource_for_user(self, user_id: int, resource_id: int) -> GeneratedResource | None:
        return next(
            (resource for resource in self.resources if resource.id == resource_id and resource.user_id == user_id),
            None,
        )

    def list_export_jobs_for_resource(self, user_id: int, resource_id: int) -> list[Any]:
        return sorted(
            [job for job in self.export_jobs if job.user_id == user_id and job.resource_id == resource_id],
            key=lambda job: (job.created_at, job.id),
            reverse=True,
        )

    def add_export_job(self, job: Any) -> Any:
        job.id = self.next_export_job_id
        self.next_export_job_id += 1
        job.created_at = NOW + timedelta(minutes=job.id)
        job.updated_at = job.created_at
        self.export_jobs.append(job)
        return job

    def get_export_job_for_user(self, user_id: int, job_id: int) -> Any | None:
        return next((job for job in self.export_jobs if job.id == job_id and job.user_id == user_id), None)

    def get_export_job(self, job_id: int) -> Any | None:
        return next((job for job in self.export_jobs if job.id == job_id), None)

    def commit(self) -> None:
        return None

    def rollback(self) -> None:
        return None

    def refresh(self, _instance: object) -> None:
        return None


def make_user(user_id: int = 1) -> User:
    return User(
        id=user_id,
        account=f"user{user_id}",
        hashed_password="not-used",
        display_name=f"学生 {user_id}",
        role="student",
        starter_mode="blank",
    )


def make_course(course_id: int = 101, owner_id: int = 1) -> Course:
    return Course(
        id=course_id,
        owner_id=owner_id,
        title="人工智能导论",
        description="课程资料",
        subject="人工智能",
        source_type="uploaded",
        visibility="private",
        status="ready",
    )


def make_point(point_id: int, title: str, order_index: int) -> KnowledgePoint:
    return KnowledgePoint(
        id=point_id,
        course_id=101,
        title=title,
        summary=f"{title} 的安全摘要。",
        chapter="第一章",
        order_index=order_index,
        difficulty=None,
        prerequisites_json=[],
    )


def make_slide_resource(resource_id: int = 9901, user_id: int = 1, course_id: int = 101) -> GeneratedResource:
    return GeneratedResource(
        id=resource_id,
        user_id=user_id,
        course_id=course_id,
        knowledge_point_id=501,
        resource_type="slide",
        title="启发式搜索个性化课件",
        agent_trace_id="trace_slide_resource",
        content_json={
            "schema_version": 2,
            "format": "rich",
            "markdown": "# 启发式搜索PPT",
            "artifact": {
                "kind": "slide_deck",
                "theme": {"name": "edunova-light", "aspect_ratio": "16:9", "accent": "#0f8f83"},
                "slides": [
                    {
                        "id": "slide-1",
                        "title": "为什么要学启发式搜索",
                        "bullets": ["减少无效搜索", "结合课程目标"],
                        "speaker_notes": "先用路径规划问题引入。",
                        "layout": "title",
                        "citation_refs": [701],
                    },
                    {
                        "id": "slide-2",
                        "title": "A* 的关键步骤",
                        "bullets": ["计算已走代价", "估计剩余代价", "选择优先节点"],
                        "speaker_notes": "结合课程例题逐步验证。",
                        "layout": "title_and_content",
                        "citation_refs": [701],
                    },
                ],
                "citation_refs": [701],
            },
            "metadata": {"review_mode": "model_and_rules"},
        },
        citation_json=[{"chunk_id": 701, "section_title": "A* 搜索", "source_title": "人工智能导论讲义.md"}],
        status="completed",
        review_status="passed",
        confidence_score=Decimal("0.88"),
        created_at=NOW,
        updated_at=NOW,
    )


def make_weakness() -> WeaknessReviewItem:
    item = WeaknessReviewItem(
        id=801,
        user_id=1,
        course_id=101,
        knowledge_point_id=402,
        title="启发式搜索",
        source_type="practice_assessment",
        status="confirmed",
        recommended_resource_ids=["901"],
        next_review_at=NOW + timedelta(days=2),
    )
    item.created_at = NOW
    item.updated_at = NOW
    return item


def make_resource() -> GeneratedResource:
    return GeneratedResource(
        id=901,
        user_id=1,
        course_id=101,
        knowledge_point_id=402,
        resource_type="doc",
        title="启发式搜索讲解",
        content_json={"markdown": "安全资源摘要", "metadata": {"generation_mode": "deterministic_source"}},
        citation_json=[],
        status="completed",
        review_status="passed",
        confidence_score=Decimal("0.86"),
        created_at=NOW,
        updated_at=NOW,
    )


def make_path_and_task() -> tuple[LearningPath, LearningTask]:
    path = LearningPath(
        id=1001,
        user_id=1,
        course_id=101,
        title="人工智能导论学习路径",
        goal="期末前掌握搜索算法",
        status="active",
        plan_json={"basis": ["课程知识点和已确认弱点"]},
        created_at=NOW,
        updated_at=NOW,
    )
    task = LearningTask(
        id=1002,
        path_id=1001,
        user_id=1,
        course_id=101,
        knowledge_point_id=402,
        title="复习启发式搜索",
        task_type="review",
        reason="来自已确认薄弱点",
        recommended_resource_ids=["901"],
        status="doing",
        due_at=NOW + timedelta(days=1),
        next_review_at=None,
        created_at=NOW,
        updated_at=NOW,
    )
    return path, task


def make_practice() -> tuple[PracticeSession, PracticeAnswer]:
    session = PracticeSession(
        id=501,
        user_id=1,
        course_id=101,
        title="人工智能导论练习",
        status="completed",
        score=Decimal("45"),
        created_at=NOW,
        updated_at=NOW,
    )
    answer = PracticeAnswer(
        id=601,
        session_id=501,
        user_id=1,
        question_json={"knowledge_point_id": "402", "knowledge_point_title": "启发式搜索", "prompt": "安全题干"},
        answer_text="这是不应进入学习档案的完整学生原始答案",
        feedback_json={"score": 20, "message": "需要复习"},
        is_correct=False,
        created_at=NOW,
    )
    return session, answer


def make_report() -> AssessmentReport:
    return AssessmentReport(
        id=701,
        user_id=1,
        course_id=101,
        practice_session_id=501,
        report_json={
            "summary": "本次评估得分 45，基于真实练习作答生成。",
            "mastery_update": {"weak_count": 1, "mastered_count": 0, "learning_count": 2},
            "weakness_list": [{"knowledge_point_id": "402", "title": "启发式搜索", "source_type": "practice_assessment"}],
            "next_step_suggestions": ["先处理启发式搜索，再回看课程引用。"],
            "evidence_refs": [{"practice_answer_id": "601", "knowledge_point_id": "402", "score": 20}],
            "review_queue_updates": [{"title": "启发式搜索", "status": "confirmed", "source_type": "practice_assessment"}],
            "profile_changes": ["不自动改写长期画像。"],
        },
        score=Decimal("45"),
        created_at=NOW,
    )


def make_repo(with_report: bool = True) -> FakeLearningDossierRepository:
    path, task = make_path_and_task()
    session, answer = make_practice()
    reports = [make_report()] if with_report else []
    return FakeLearningDossierRepository(
        courses=[make_course(), make_course(202, owner_id=2)],
        knowledge_points=[make_point(401, "人工智能概述", 0), make_point(402, "启发式搜索", 1)],
        weakness_items=[make_weakness()],
        paths=[path],
        tasks=[task],
        resources=[make_resource()],
        practice_sessions=[session],
        practice_answers=[answer],
        reports=reports,
    )


def as_dict(model: Any) -> dict[str, Any]:
    return model.model_dump() if hasattr(model, "model_dump") else model


def make_export_settings(tmp_path: Path) -> Settings:
    return Settings(
        _env_file=None,
        jwt_secret="export-test-secret-with-32-bytes-long",
        jwt_expire_minutes=30,
        export_dir=str(tmp_path / "exports"),
    )


def test_learning_dossier_exports_ready_report_path_resources_and_safe_markdown() -> None:
    from backend.app.services.exports import ExportService

    service = ExportService(make_repo(with_report=True))

    result = as_dict(service.export_learning_dossier(make_user(), course_id=101))

    assert result["course_id"] == "101"
    assert result["content_type"] == "text/markdown; charset=utf-8"
    assert result["filename"].endswith(".md")
    assert "人工智能导论" in result["filename"]
    assert result["source_summary"]["has_report"] is True
    assert result["source_summary"]["report_id"] == "701"
    assert result["source_summary"]["knowledge_point_count"] == 2
    assert result["source_summary"]["weakness_count"] == 1
    assert result["source_summary"]["path_task_count"] == 1
    assert result["source_summary"]["resource_count"] == 1
    assert result["source_summary"]["practice_answer_count"] == 1
    markdown = result["markdown"]
    assert "# 人工智能导论 学习档案" in markdown
    assert "本次评估得分 45" in markdown
    assert "启发式搜索" in markdown
    assert "复习启发式搜索" in markdown
    assert "启发式搜索讲解" in markdown
    assert "到期" not in markdown
    assert "这是不应进入学习档案的完整学生原始答案" not in markdown
    assert "系统提示词" not in markdown
    assert "模型输入" not in markdown
    assert "API Key" not in markdown
    assert "JWT" not in markdown


def test_learning_dossier_exports_honest_empty_report_state() -> None:
    from backend.app.services.exports import ExportService

    service = ExportService(make_repo(with_report=False))

    result = as_dict(service.export_learning_dossier(make_user(), course_id=101))

    assert result["source_summary"]["has_report"] is False
    assert result["source_summary"]["report_id"] is None
    assert "还没有真实学习报告。" in result["markdown"]
    assert "完成一次课程练习后生成报告。" in result["markdown"]
    assert "本次评估得分 45" not in result["markdown"]


def test_learning_dossier_validates_course_scope() -> None:
    from backend.app.services.exports import ExportNotFoundError, ExportService

    service = ExportService(make_repo(with_report=True))

    with pytest.raises(ExportNotFoundError):
        service.export_learning_dossier(make_user(), course_id=202)


def test_learning_dossier_route_requires_login_and_returns_envelope() -> None:
    from backend.app.api.v1.exports import get_export_service
    from backend.app.services.exports import ExportService

    repo = make_repo(with_report=True)
    user = make_user()
    settings = Settings(_env_file=None, jwt_secret="export-test-secret-with-32-bytes-long", jwt_expire_minutes=30)
    app = create_app()
    app.dependency_overrides[get_auth_service] = lambda: AuthService(repository=TokenAuthRepository(user), settings=settings)
    app.dependency_overrides[get_export_service] = lambda: ExportService(repo)
    client = TestClient(app)
    token = create_access_token(str(user.id), settings=settings)
    headers = {"Authorization": f"Bearer {token}"}

    unauthorized = client.post("/api/v1/exports/learning-dossier", json={"course_id": 101})
    exported = client.post("/api/v1/exports/learning-dossier", headers=headers, json={"course_id": 101})
    missing = client.post("/api/v1/exports/learning-dossier", headers=headers, json={"course_id": 202})

    assert unauthorized.status_code == 401
    assert exported.status_code == 200
    assert exported.json()["data"]["filename"].endswith(".md")
    assert "人工智能导论 学习档案" in exported.json()["data"]["markdown"]
    assert missing.status_code == 404


@pytest.mark.parametrize(
    ("export_format", "expected_content_type", "magic"),
    [
        ("markdown", "text/markdown; charset=utf-8", b"# "),
        ("pdf", "application/pdf", b"%PDF"),
        ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", b"PK"),
    ],
)
def test_learning_dossier_export_jobs_generate_files(
    tmp_path: Path,
    export_format: str,
    expected_content_type: str,
    magic: bytes,
) -> None:
    from backend.app.services.exports import ExportService

    repo = make_repo(with_report=True)
    service = ExportService(repo, settings=make_export_settings(tmp_path), run_jobs_inline=True)

    job = as_dict(service.create_learning_dossier_job(make_user(), course_id=101, export_format=export_format))

    assert job["job_id"] == "1"
    assert job["status"] == "completed"
    assert job["format"] == export_format
    assert job["content_type"] == expected_content_type
    assert job["agent_trace_id"]
    assert job["error_message"] is None
    saved_path = Path(repo.export_jobs[0].file_path)
    assert saved_path.exists()
    assert saved_path.read_bytes().startswith(magic)
    assert "系统提示词" not in saved_path.read_bytes().decode("utf-8", errors="ignore")


def test_learning_dossier_pdf_falls_back_when_cjk_font_is_not_tt_compatible(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from backend.app.services.exports import ExportService

    service = ExportService(make_repo(with_report=True), settings=make_export_settings(tmp_path))
    incompatible_font = tmp_path / "unsupported-cjk-font.ttc"
    incompatible_font.write_bytes(b"not a real TrueType font")
    monkeypatch.setattr(service, "_find_cjk_font", lambda: incompatible_font)

    rendered = service._render_pdf("# 人工智能导论 学习档案\n\n- 中文可读测试\n")

    assert rendered.startswith(b"%PDF")


def test_learning_dossier_export_job_rejects_other_users_job(tmp_path: Path) -> None:
    from backend.app.services.exports import ExportNotFoundError, ExportService

    service = ExportService(make_repo(with_report=True), settings=make_export_settings(tmp_path), run_jobs_inline=True)
    service.create_learning_dossier_job(make_user(1), course_id=101, export_format="markdown")

    with pytest.raises(ExportNotFoundError):
        service.get_export_job(make_user(2), job_id=1)


def test_learning_dossier_export_job_route_creates_reads_and_downloads(tmp_path: Path) -> None:
    from backend.app.api.v1.exports import get_export_service
    from backend.app.services.exports import ExportService

    repo = make_repo(with_report=True)
    user = make_user()
    settings = make_export_settings(tmp_path)
    app = create_app()
    app.dependency_overrides[get_auth_service] = lambda: AuthService(repository=TokenAuthRepository(user), settings=settings)
    app.dependency_overrides[get_export_service] = lambda: ExportService(repo, settings=settings, run_jobs_inline=True)
    client = TestClient(app)
    token = create_access_token(str(user.id), settings=settings)
    headers = {"Authorization": f"Bearer {token}"}

    created = client.post(
        "/api/v1/exports/learning-dossier/jobs",
        headers=headers,
        json={"course_id": 101, "format": "pdf"},
    )
    job_id = created.json()["data"]["job_id"]
    status_response = client.get(f"/api/v1/exports/{job_id}", headers=headers)
    download_response = client.get(f"/api/v1/exports/{job_id}/download", headers=headers)

    assert created.status_code == 200
    assert created.json()["data"]["status"] == "completed"
    assert status_response.status_code == 200
    assert status_response.json()["data"]["format"] == "pdf"
    assert download_response.status_code == 200
    assert download_response.headers["content-type"].startswith("application/pdf")
    assert download_response.content.startswith(b"%PDF")


def test_resource_slide_export_job_generates_real_pptx_and_reuses_completed_job(tmp_path: Path) -> None:
    from pptx import Presentation

    from backend.app.services.exports import ExportService

    repo = make_repo(with_report=True)
    repo.resources.append(make_slide_resource())
    service = ExportService(repo, settings=make_export_settings(tmp_path), run_jobs_inline=True)

    first = as_dict(service.create_resource_export_job(make_user(), resource_id=9901, export_format="pptx"))
    second = as_dict(service.create_resource_export_job(make_user(), resource_id=9901, export_format="pptx"))

    assert first["status"] == "completed"
    assert first["format"] == "pptx"
    assert first["export_type"] == "resource_artifact"
    assert first["resource_id"] == "9901"
    assert first["job_id"] == second["job_id"]
    saved_path = Path(repo.export_jobs[0].file_path)
    assert saved_path.exists()
    assert saved_path.read_bytes().startswith(b"PK")
    presentation = Presentation(saved_path)
    assert len(presentation.slides) == 3
    slide_text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text"))
    assert "为什么要学启发式搜索" in slide_text
    assert "课程引用" in slide_text


def test_resource_export_rejects_non_slide_and_other_user(tmp_path: Path) -> None:
    from backend.app.services.exports import ExportNotFoundError, ExportService, ExportValidationError

    repo = make_repo(with_report=True)
    repo.resources.extend(
        [
            make_slide_resource(),
            GeneratedResource(
                id=902,
                user_id=1,
                course_id=101,
                resource_type="doc",
                title="讲解",
                content_json={"markdown": "讲解"},
                citation_json=[],
                status="completed",
                review_status="passed",
            ),
        ]
    )
    service = ExportService(repo, settings=make_export_settings(tmp_path), run_jobs_inline=True)

    with pytest.raises(ExportValidationError):
        service.create_resource_export_job(make_user(), resource_id=902, export_format="pptx")
    with pytest.raises(ExportNotFoundError):
        service.create_resource_export_job(make_user(2), resource_id=9901, export_format="pptx")


def test_resource_export_routes_create_list_and_download_pptx(tmp_path: Path) -> None:
    from backend.app.api.v1.exports import get_export_service
    from backend.app.api.v1.resources import get_resource_export_service
    from backend.app.services.exports import ExportService

    repo = make_repo(with_report=True)
    repo.resources.append(make_slide_resource())
    user = make_user()
    settings = make_export_settings(tmp_path)
    service = ExportService(repo, settings=settings, run_jobs_inline=True)
    app = create_app()
    app.dependency_overrides[get_auth_service] = lambda: AuthService(repository=TokenAuthRepository(user), settings=settings)
    app.dependency_overrides[get_resource_export_service] = lambda: service
    app.dependency_overrides[get_export_service] = lambda: service
    client = TestClient(app)
    token = create_access_token(str(user.id), settings=settings)
    headers = {"Authorization": f"Bearer {token}"}

    created = client.post("/api/v1/resources/9901/exports", headers=headers, json={"format": "pptx"})
    listed = client.get("/api/v1/resources/9901/exports", headers=headers)
    job_id = created.json()["data"]["job_id"]
    downloaded = client.get(f"/api/v1/exports/{job_id}/download", headers=headers)

    assert created.status_code == 200
    assert created.json()["data"]["status"] == "completed"
    assert listed.status_code == 200
    assert listed.json()["data"][0]["resource_id"] == "9901"
    assert downloaded.status_code == 200
    assert downloaded.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert downloaded.content.startswith(b"PK")
